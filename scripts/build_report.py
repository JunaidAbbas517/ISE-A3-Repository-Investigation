"""Build editable study report and exact engineering diagrams.

Authoring dependencies are separate from the investigated application's lock.
The distributed DOCX is the editable report source; PDF is rendered from it.
"""
from pathlib import Path
import json
import textwrap
import matplotlib
matplotlib.use('Agg')
import matplotlib.pyplot as plt
from matplotlib.patches import FancyBboxPatch, FancyArrowPatch
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

R=Path(__file__).resolve().parents[1]
meta=json.loads((R/'upstream/UPSTREAM.json').read_text())
summary=json.loads((R/'evidence/verified/summary.json').read_text())
risks=json.loads((R/'analysis/risks.json').read_text())

def diagram(name,title,nodes,edges,size):
 plt.rcParams.update({'font.family':'DejaVu Sans','svg.fonttype':'none'})
 fig,ax=plt.subplots(figsize=size);fig.patch.set_facecolor('white')
 ax.set(xlim=(0,1),ylim=(0,1));ax.axis('off')
 ax.text(.02,.98,title,fontsize=15,fontweight='bold',va='top',color='#152d40')
 for ident,(x,y,w,h,text,fill) in nodes.items():
  ax.add_patch(FancyBboxPatch((x,y),w,h,boxstyle='round,pad=0.009,rounding_size=0.012',facecolor=fill,edgecolor='#8d9daa',linewidth=1))
  ax.text(x+w/2,y+h/2,text,ha='center',va='center',fontsize=12.5,color='#172b3a',linespacing=1.3)
 for start,end,mode in edges:
  a=nodes[start];b=nodes[end]
  if mode=='down': p=(a[0]+a[2]/2,a[1]);q=(b[0]+b[2]/2,b[1]+b[3])
  elif mode=='right':p=(a[0]+a[2],a[1]+a[3]/2);q=(b[0],b[1]+b[3]/2)
  else:p=(a[0]+a[2]/2,a[1]);q=(b[0]+b[2]/2,b[1]+b[3])
  ax.add_patch(FancyArrowPatch(p,q,arrowstyle='-|>',mutation_scale=14,linewidth=1.3,color='#567084',connectionstyle='arc3,rad=0'))
 fig.subplots_adjust(0,0,1,1)
 for suffix in ['svg','png']:fig.savefig(R/'diagrams'/f'{name}.{suffix}',dpi=200,facecolor='white')
 plt.close(fig)

diagram('architecture','Observed local architecture',{
 'cli':(.05,.76,.90,.13,'CLI entry and commands\nClick arguments, configuration and display','#e8eff7'),
 'config':(.05,.46,.30,.16,'Trusted local config\nExecutable Python','#fff4df'),
 'domain':(.44,.46,.51,.16,'Task model and Database\nDue values, UID and SEQUENCE','#e7f2ed'),
 'files':(.44,.12,.23,.17,'VTODO files\nDurable .ics data','#e8eff7'),
 'cache':(.72,.12,.23,.17,'SQLite cache\nQuery data and IDs','#f0edf6')
},[('cli','config','down'),('cli','domain','down'),('domain','files','down'),('domain','cache','down')],(9,4.6))

diagram('call-path','Observed deadline edit path',{
 'a':(.09,.81,.82,.105,'1  CLI startup and trusted configuration','#e8eff7'),
 'b':(.09,.665,.82,.105,'2  Parse due value and load the existing Todo','#e8eff7'),
 'c':(.09,.52,.82,.105,'3  Database.save increments SEQUENCE','#e7f2ed'),
 'd':(.09,.375,.82,.105,'4  VtodoWriter serializes the calendar record','#e7f2ed'),
 'e':(.09,.23,.82,.105,'5  Atomic replacement writes the .ics file','#e8eff7'),
 'f':(.09,.085,.82,.105,'6  Refresh SQLite cache and display the result','#f0edf6')
},[('a','b','down'),('b','c','down'),('c','d','down'),('d','e','down'),('e','f','down')],(9,5.2))

diagram('change-boundary','Proposed future boundary  Not implemented in A3',{
 'input':(.04,.50,.24,.23,'Notice and\nowner decision','#e8eff7'),
 'service':(.38,.50,.25,.23,'A2 verification\nand protected history','#e7f2ed'),
 'export':(.73,.50,.23,.23,'Controlled export\nto personal calendar','#e8eff7'),
 'outside':(.38,.09,.25,.19,'Official LMS deadline\nremains external','#fff4df')
},[('input','service','right'),('service','export','right')],(9,3.2))

d=Document();sec=d.sections[0]
sec.page_width=Inches(8.27);sec.page_height=Inches(11.69)
sec.top_margin=Inches(.70);sec.bottom_margin=Inches(.65)
sec.left_margin=sec.right_margin=Inches(.80)
normal=d.styles['Normal'];normal.font.name='Arial';normal.font.size=Pt(10.5)
normal.paragraph_format.space_after=Pt(6);normal.paragraph_format.line_spacing=1.08
for name,size in [('Title',19),('Subtitle',11),('Heading 1',14),('Heading 2',11)]:
 s=d.styles[name];s.font.name='Arial';s.font.size=Pt(size);s.font.color.rgb=RGBColor(0,0,0)
 s.paragraph_format.space_before=Pt(10 if name.startswith('Heading') else 0)
 s.paragraph_format.space_after=Pt(6)
 if name.startswith('Heading'):s.paragraph_format.keep_with_next=True

footer=sec.footer.paragraphs[0];footer.alignment=WD_ALIGN_PARAGRAPH.RIGHT
r=footer.add_run('ISE A3  |  ');r.font.size=Pt(8);r.font.color.rgb=RGBColor.from_string('667788')
field=OxmlElement('w:fldSimple');field.set(qn('w:instr'),'PAGE');footer._p.append(field)

def p(text,style=None):
 previous=list(d.element.body)[-2] if len(d.element.body)>1 else None
 paragraph=d.add_paragraph(text,style)
 if previous is not None and previous.tag == qn('w:tbl'):
  paragraph.paragraph_format.space_before=Pt(5)
 return paragraph
def h(text):return d.add_heading(text,level=1)
def sub(text):return d.add_heading(text,level=2)
def page(title):d.add_page_break();h(title)
def image(name,width=6.6):d.add_picture(str(R/'diagrams'/f'{name}.png'),width=Inches(width))
def table(headers,rows,widths):
 t=d.add_table(rows=1,cols=len(headers));t.alignment=WD_TABLE_ALIGNMENT.CENTER;t.autofit=False
 for col,w in zip(t.columns,widths):col.width=Inches(w)
 for c,txt in zip(t.rows[0].cells,headers):c.text=txt
 for row in rows:
  for c,txt in zip(t.add_row().cells,row):c.text=str(txt)
 for i,row in enumerate(t.rows):
  trPr=row._tr.get_or_add_trPr();cant=OxmlElement('w:cantSplit');trPr.append(cant)
  if i==0:
   repeat=OxmlElement('w:tblHeader');trPr.append(repeat)
  for j,c in enumerate(row.cells):
   c.width=Inches(widths[j]);c.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
   pr=c._tc.get_or_add_tcPr();shade=OxmlElement('w:shd');shade.set(qn('w:fill'),'173E59' if i==0 else ('F0F4F7' if i%2 else 'FFFFFF'));pr.append(shade)
   borders=OxmlElement('w:tcBorders')
   for edge in ['top','left','bottom','right']:
    el=OxmlElement('w:'+edge);el.set(qn('w:val'),'single');el.set(qn('w:sz'),'4');el.set(qn('w:color'),'D9D9D9');borders.append(el)
   pr.append(borders)
   margins=OxmlElement('w:tcMar')
   for edge in ['top','left','bottom','right']:
    el=OxmlElement('w:'+edge);el.set(qn('w:w'),'75');el.set(qn('w:type'),'dxa');margins.append(el)
   pr.append(margins)
   for par in c.paragraphs:
    par.paragraph_format.space_after=Pt(2);par.paragraph_format.space_before=Pt(2)
    par.paragraph_format.line_spacing=1.03
    for run in par.runs:
     run.font.size=Pt(9.2);run.font.color.rgb=RGBColor.from_string('FFFFFF' if i==0 else '172B3A');run.bold=i==0
 return t
def code(text):
 for line in text.strip().splitlines():
  par=p(line);par.paragraph_format.space_after=Pt(2)
  for run in par.runs:run.font.name='Courier New';run.font.size=Pt(8.5)

d.add_paragraph('Repository Investigation and Reproducible Build','Title')
p('ISE A3  |  Intelligent Software Engineering  |  25 October 2026','Subtitle')
p('Project context: Assignment Deadline Change Tracking and Verification System')
h('Purpose and principal finding')
p('This investigation examines Todoman, an unfamiliar personal task repository, as a possible storage and presentation boundary for the deadline-change workflow established in A1 and specified in A2. Ordinary due-date editing is reproducible. Academic verification, stale-version protection and explicit-instant semantics require a controlled boundary beyond the examined CLI write contract.')
p('The study freezes an upstream commit, builds its source in an isolated environment, reconstructs one executed deadline edit, and compares observed behavior with A2. It retains negative results and distinguishes facts, inferences and proposed changes. Upstream source remains unchanged.')
h('Frozen repository and scope')
table(['Item','Recorded identity'],[
 ('Repository',meta['repository_url']),('Commit',meta['commit']),
 ('Snapshot label',meta['snapshot_version']+'; study label, not a released version'),
 ('Licence and input','ISC; complete source export and synthetic VTODO fixture retained'),
 ('Execution target','CPython 3.12.14; Linux x86_64; UTC; offline locked wheels')],[1.25,5.35])
h('Recorded outcomes')
table(['Evidence','Outcome','Interpretation'],[
 ('Upstream cases','197 total','188 passed; 6 skipped; 3 expected failures'),
 ('Study checks','6 observations passed','Includes observed missing safeguards'),
 ('Clean build and start','Successful','Wheel built, installed and CLI started offline'),
 ('Future integration','Proposed only','Verification service and export adapter are not implemented')],[1.40,1.45,3.75])
p('All task records are synthetic. No real student information, live university source, institutional credentials or LMS connection was used. A1 simulated pilot data remain simulated.',None)

page('Reproduction method and environment')
p('The complete offline procedure is scripts/reproduce.py. It verifies 33 input hashes, creates a new virtual environment, installs 30 pinned dependency distributions, builds and installs the exported source, executes tests, starts the CLI, runs the probes and checks dependencies. Each run writes a new directory and retains stage commands, exit codes and raw output.')
code('python3 --version\npython3 scripts/reproduce.py --verify-only\npython3 scripts/reproduce.py --output runs/local')
sub('Environment and build identity')
p('The runtime record identifies CPython 3.12.14, Linux x86_64, glibc 2.39 and the SQLite version. TZ=UTC, LC_ALL=C.UTF-8, PYTHONHASHSEED=0, DETERMINISTIC_TESTS=true and Hypothesis seed 20261025 are fixed. Each dependency wheel is hash-locked. Native Windows is excluded by upstream; WSL/Linux guidance is provided but WSL was not separately tested.')
p('A Git export contains no SCM metadata. The build therefore supplies SETUPTOOLS_SCM_PRETEND_VERSION_FOR_TODOMAN with the documented snapshot label. SOURCE_DATE_EPOCH is derived from the pinned commit timestamp. Two clean builds produced the same recorded wheel hash. This does not establish reproducibility on every platform.')
sub('Retained test exclusions')
table(['Cases','Reason','Treatment'],[
 ('1 skip','Optional click_repl is absent','Optional REPL is not tested'),
 ('4 skips','Invalid recurrence combinations','Preserved as upstream exclusions'),
 ('1 skip','UI case linked to issue 537','No interactive UI guarantee'),
 ('3 expected failures','Issue 9 and two unimplemented config cases','Retained, never counted as passes')],[.8,2.8,3.0])
p('The suite reports 16 warnings, principally dependency API deprecations. Default coverage flags are explicitly disabled; no coverage percentage is claimed. Raw XML and text preserve every test and reason.')
sub('Failed and nondeterministic evidence')
p('The first study harness failed because an icalendar list was passed directly to next(). That investigator error is retained and corrected with iter(...). Malformed input and missing configuration intentionally produce nonzero exit codes. Absolute due values and UID are fixed; LAST-MODIFIED, mtimes, paths and timings vary and are not byte-identity oracles.')

page('Architecture and dependency boundaries')
image('architecture')
p('Figure 1. Observed local architecture. Lines represent examined dependencies and calls; no external network synchronization was executed.')
table(['Component','Source location','Role'],[
 ('Entry and commands','__main__.py; cli.py','CLI startup, validation and edit/show/list commands'),
 ('Configuration','configuration.py:220','Executes trusted local config and validates settings'),
 ('Dates and feedback','formatters.py:255,262','Parses configured local dates/times and formats output'),
 ('Persistence','model.py:1064,1164','Database facade writes VTODO and refreshes cache'),
 ('Files and cache','model.py:296,422','Durable ICS records and derived SQLite query data')],[1.30,1.90,3.40])
p('Critical dependencies are Click/click-log, icalendar, python-dateutil, parsedatetime, pyxdg, humanize and urwid. Build and test tools and all transitive installed distributions are locked separately from standard-library sqlite3. Urwid widgets were inspected but not interactively exercised.')
p('Interfaces include CLI arguments, JSON porcelain output, local files and optional editor integration. Separate synchronization software may consume calendar files; no live CalDAV or LMS integration was tested. SQLite also assigns local task IDs, which must not be assumed to be institutional assignment identifiers.')

page('Executed deadline edit and evidence chain')
image('call-path')
p('Figure 2. Reconstructed path supported by the selected runtime trace, source definitions and persisted data.')
p('P1 begins with a fixed synthetic UID, DUE 10 October 2026 at 16:30 UTC and SEQUENCE 1. An edit proposes 17 October at 16:30 UTC. The resulting ICS file stores that instant and SEQUENCE 2; a fresh CLI process and direct SQLite observation agree. This UTC fixture is separate from the Asia/Shanghai A2 examples.')
p('The profiler wrapper invokes the real todoman.__main__ module without changing application source. Its 28 selected calls include configuration, date validation, task retrieval, Database.save, VtodoWriter, atomic_write, cache refresh/commit and feedback. It is a filtered call trace, not a complete coverage result.')
table(['Classification','Supported statement'],[
 ('Runtime fact','Normal edit, file replacement, cache row and fresh-process readback were observed.'),
 ('Source fact','File replacement precedes a separate SQLite cache commit.'),
 ('Inference','An interruption between those operations could create temporary divergence; not fault-injected.'),
 ('Proposal','A separate service would own verified academic records and export approved copies.')],[1.20,5.40])

page('Characterization checks and counterexamples')
p('The six checks inspect files, cache rows, actual exit codes or a harmless execution marker. A passing check means the stated behavior was observed, including behaviors that fail the A2 contract. It does not establish academic or production security.')
table(['ID','Procedure','Observed result and limit'],[
 ('P1','Normal deadline edit','Due value persists; SEQUENCE increases from 1 to 2; a new CLI process and SQLite cache agree. No institutional source or owner verification is established.'),
 ('P2','Malformed date','Unrecognized input exits 2; the existing ICS file remains byte-identical. This covers one malformed date, not an exhaustive parser study.'),
 ('P3','Stale object saves','Two Todo objects are read before either save. Saving both accepts the stale replacement and repeats the resulting sequence. This is a deterministic interleaving, not a measured parallel race.'),
 ('P4','Date-only deadline','The CLI accepts 25 October 2026 and persists an all-day date. That is legitimate task behavior but lacks the exact cutoff instant required by A2.'),
 ('P5','Executable configuration','Loading a synthetic Python config writes an isolated harmless marker. This confirms that config belongs within a trusted local deployment boundary.'),
 ('P6','Missing configuration','An explicitly nonexistent config exits 1 with actionable feedback. No real account or academic record is involved.')],[.45,1.30,4.85])
sub('Evidence retained')
p('evidence/verified/workflow contains command stdout/stderr, results.json, call-trace.json, original and revised calendar snapshots, and cache-after-edit.json. The upstream JUnit XML preserves expected failures and exclusions. The failed first harness attempt remains separately identified as an investigator error.')
sub('Relationship to the earlier assignments')
p('A1 defines the synchronization problem; its simulated pilot does not prove prevalence. A2 defines the desired owner, source, version, rejection and time rules. Todoman has a different personal-task contract. The observed gaps identify an adaptation boundary rather than automatically constituting defects in upstream software.')

page('Three risks and acceptance conditions')
p('Risk priority is qualitative and scoped to reusing the repository within the academic workflow. No probability, incident frequency or CVSS score is fabricated.')
for risk in risks:
 sub(risk['id']+' '+risk['risk'])
 p(risk['condition']+' '+risk['evidence']+' '+risk['impact'])
 p('Response: '+risk['response']+' Acceptance: '+risk['acceptance'])
 p('Residual limit: '+risk['residual'])
p('Owners are future implementation/deployment roles, not named reviewers or assigned individuals. The complete register, including conditions, priorities, ownership and residual limits, is supplied in Markdown, CSV and JSON.')

page('Architecture decision and alternatives')
p('ADR 001 addresses R1: how to add the A2 verification and version boundary while preserving a reproducible upstream baseline. Status: proposed for future implementation; A3 makes no product change.')
table(['Criterion','A Modify upstream write core','B Separate service and export'],[
 ('Rule ownership','Add academic concepts to CLI/model','Keep A2 rules in a protected service'),
 ('Change surface','Reconcile command, model, cache and other writers','Add a bounded service and export adapter'),
 ('Version safety','Retrofit coordinated conflict checks/history','Transactional service record; export after commit'),
 ('Existing writers','Every write path must adopt the rules','External files remain derived, unverified copies'),
 ('Main cost','Fork maintenance and compatibility','New identity mapping, store and recovery behavior')],[1.05,2.775,2.775])
sub('Proposed decision')
p('Prefer B. A future service would authenticate ownership, verify source authority independently, validate candidate state and explicit time semantics, compare the expected version and commit one history entry before export. It would own the personal verified record, not the official university deadline. Todoman would be an optional display or export target.')
p('The service would map stable assignment identity to a VTODO UID rather than treat the cache ID as an institutional identifier. Uploaded documents would be processed as data and never loaded as Python configuration. External edits would become new unverified input rather than silently replacing the verified service record.')
sub('Consequences and reversal')
p('The selected boundary adds authentication, storage and export reconciliation work. It cannot constrain arbitrary writers of exported calendars; protection depends on the service retaining authority over its own record. Future acceptance must include non-owner, unverified-source, stale-version, retry, rejection and time-zone cases. No such implementation is claimed here.')
p('For A3, rollback means discarding generated workspaces; upstream source is unchanged. A future deployment must be able to disable export and restore its prior service release without losing revision history. No production migration or rollback was executed.')

page('Safe change boundary and review handoff')
image('change-boundary')
p('Figure 3. Proposed boundary for a later bounded implementation. The official LMS remains authoritative outside the personal verification service.')
sub('What the evidence supports')
p('The recorded Linux environment can install its locked dependencies offline, build the frozen source, execute the upstream suite and complete the six isolated observations. The call-path reconstruction is supported by source and selected runtime calls. This establishes a reviewable investigation baseline, not production authorization, source authenticity or operational reliability.')
sub('Limits and unresolved questions')
p('No live LMS/CalDAV, real participant evidence, interactive usability session, native Windows/WSL execution, DST study, process-crash recovery test or production deployment was performed. The file/cache interruption window remains an inference. The optional REPL and upstream exclusions remain unverified. A future implementation must establish a credible institutional source-authority procedure and confirm user value with authentic evidence.')
sub('Provenance and AI use')
p('Primary source: '+meta['commit_url'])
p('The unchanged source archive, licence, file hashes, runtime lock, raw output and decision record are included. ChatGPT/Codex assisted with discovery, execution, analysis and writing. AI-USE-AND-VERIFICATION.md records accepted/rejected advice and programmatic verification. No independent human or student rerun is claimed.')
sub('Submission handoff')
p('Upload the extracted package contents to the student study repository, review the evidence, and create a stable commit/tag. Submit that URL with the 250-word summary in PLATFORM-SUBMISSION.md. The upstream commit URL identifies the investigated repository and does not replace the student evidence repository URL.')

for tree in [d.styles.element, d.element]:
 for element in list(tree.iter(qn('w:pBdr'))):
  element.getparent().remove(element)
out=R/'report/ISE-A3-Final-Report.docx';d.save(out)
print(out)
